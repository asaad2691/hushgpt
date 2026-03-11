import csv
import io
import json
from pathlib import Path
from uuid import uuid4

from app.services.llm_provider import get_client


class FileToolsService:
    SUPPORTED_FORMATS = {"txt", "md", "csv", "html", "docx", "pdf", "json"}
    SUPPORTED_PARSERS = {"table", "resume", "invoice", "contract", "section-summary"}

    def __init__(self, app_config, output_dir):
        self.max_new_tokens = int(app_config.get("FILE_MAX_NEW_TOKENS", 1200))
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def analyze(self, file_bytes, filename, prompt):
        text = self._extract_text(file_bytes, filename)
        ask = prompt or "Summarize this file and list key points."
        chunks = self._chunk_text(text, chunk_size=5000, overlap=400)
        chunk_notes = []
        for idx, chunk in enumerate(chunks[:4], start=1):
            llm_prompt = (
                f"Filename: {filename}\n"
                f"User request: {ask}\n"
                f"Chunk {idx} of {min(len(chunks), 4)}:\n"
                f"{chunk}\n\n"
                "Extract only the important points relevant to the user's request."
            )
            result = get_client().chat_once(
                messages=[{"role": "user", "content": llm_prompt}],
                options={"max_new_tokens": 220, "temperature": 0.2, "do_sample": False},
            )
            note = (result.get("message", {}).get("content", "") or "").strip()
            if note:
                chunk_notes.append(f"Chunk {idx} notes:\n{note}")

        synthesis_prompt = (
            f"Filename: {filename}\n"
            f"User request: {ask}\n\n"
            "Combine the extracted notes below into one direct, useful answer.\n"
            "Use bullets when appropriate and keep the answer grounded in the file.\n\n"
            + "\n\n".join(chunk_notes or ["No chunk notes available."])
        )
        result = get_client().chat_once(
            messages=[{"role": "user", "content": synthesis_prompt}],
            options={"max_new_tokens": self.max_new_tokens, "temperature": 0.2, "do_sample": False},
        )
        return {
            "analysis": result.get("message", {}).get("content", ""),
            "text_preview": text[:1200],
            "full_text": text,
        }

    def generate(self, prompt, target_format, filename=None):
        target_format = self._normalize_format(target_format)
        llm_prompt = (
            f"Generate content for a {target_format.upper()} file.\n"
            f"User request: {prompt}\n"
            "Return clean content only."
        )
        result = get_client().chat_once(
            messages=[{"role": "user", "content": llm_prompt}],
            options={"max_new_tokens": self.max_new_tokens},
        )
        text = (result.get("message", {}).get("content", "") or "").strip()
        return self._write_output(text, target_format, filename=filename)

    def convert(self, file_bytes, filename, target_format):
        target_format = self._normalize_format(target_format)
        text = self._extract_text(file_bytes, filename)
        return self._write_output(text, target_format)

    def compare(self, left_bytes, left_filename, right_bytes, right_filename, prompt=""):
        left_text = self._extract_text(left_bytes, left_filename)
        right_text = self._extract_text(right_bytes, right_filename)
        ask = prompt or "Compare these two files and highlight important differences, overlaps, and risks."
        llm_prompt = (
            f"Left file: {left_filename}\n{left_text[:9000]}\n\n"
            f"Right file: {right_filename}\n{right_text[:9000]}\n\n"
            f"User request: {ask}\n"
            "Return a practical comparison with similarities, differences, and recommended next actions."
        )
        result = get_client().chat_once(
            messages=[{"role": "user", "content": llm_prompt}],
            options={"max_new_tokens": self.max_new_tokens, "temperature": 0.2, "do_sample": False},
        )
        return {
            "comparison": result.get("message", {}).get("content", ""),
            "left_preview": left_text[:800],
            "right_preview": right_text[:800],
            "left_full_text": left_text,
            "right_full_text": right_text,
        }

    def parse(self, file_bytes, filename, parser_type):
        parser_type = (parser_type or "").strip().lower()
        if parser_type not in self.SUPPORTED_PARSERS:
            raise RuntimeError(f"Unsupported parser_type: {parser_type}")
        text = self._extract_text(file_bytes, filename)
        if parser_type == "table":
            return self._parse_table(text, filename)
        if parser_type == "resume":
            return self._parse_resume(text, filename)
        if parser_type == "invoice":
            return self._parse_invoice(text, filename)
        if parser_type == "contract":
            return self._parse_contract(text, filename)
        return self._parse_sections(text, filename)

    def _normalize_format(self, fmt):
        fmt = (fmt or "").strip().lower().lstrip(".")
        if fmt not in self.SUPPORTED_FORMATS:
            raise RuntimeError(f"Unsupported format: {fmt}")
        return fmt

    def _extract_text(self, file_bytes, filename):
        ext = Path(filename or "file.txt").suffix.lower().lstrip(".")
        data = file_bytes.decode("utf-8", errors="ignore")

        if ext in {"txt", "md", "html", "htm", "json"}:
            if ext == "json":
                try:
                    parsed = json.loads(data)
                    return json.dumps(parsed, indent=2, ensure_ascii=False)
                except json.JSONDecodeError:
                    return data
            return data

        if ext == "csv":
            reader = csv.reader(io.StringIO(data))
            rows = [" | ".join(row) for row in reader]
            return "\n".join(rows)

        if ext == "docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("python-docx is required. Install: pip install python-docx") from exc
            doc = Document(io.BytesIO(file_bytes))
            return "\n".join(p.text for p in doc.paragraphs if p.text)

        if ext == "pdf":
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError("pypdf is required. Install: pip install pypdf") from exc
            reader = PdfReader(io.BytesIO(file_bytes))
            return "\n".join((page.extract_text() or "") for page in reader.pages)

        return data

    @staticmethod
    def _chunk_text(text, chunk_size=5000, overlap=400):
        text = (text or "").strip()
        if not text:
            return [""]

        chunks = []
        start = 0
        length = len(text)
        while start < length:
            end = min(length, start + chunk_size)
            chunks.append(text[start:end])
            if end >= length:
                break
            start = max(end - overlap, start + 1)
        return chunks

    def _write_output(self, text, target_format, filename=None):
        safe_name = (filename or f"file_{uuid4().hex[:8]}").strip()
        safe_name = safe_name.replace("/", "_").replace("\\", "_")
        output_path = self.output_dir / f"{safe_name}.{target_format}"

        if target_format in {"txt", "md", "html"}:
            output_path.write_text(text, encoding="utf-8")
        elif target_format == "json":
            try:
                parsed = json.loads(text)
            except json.JSONDecodeError:
                parsed = {"content": text}
            output_path.write_text(json.dumps(parsed, indent=2, ensure_ascii=False), encoding="utf-8")
        elif target_format == "csv":
            with output_path.open("w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                for line in text.splitlines():
                    writer.writerow([line])
        elif target_format == "docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("python-docx is required. Install: pip install python-docx") from exc
            doc = Document()
            for line in text.splitlines() or [text]:
                doc.add_paragraph(line)
            doc.save(str(output_path))
        elif target_format == "pdf":
            try:
                from reportlab.pdfgen import canvas
            except ImportError as exc:
                raise RuntimeError("reportlab is required. Install: pip install reportlab") from exc
            c = canvas.Canvas(str(output_path))
            y = 800
            for line in text.splitlines() or [text]:
                c.drawString(40, y, line[:120])
                y -= 16
                if y < 40:
                    c.showPage()
                    y = 800
            c.save()
        else:
            raise RuntimeError(f"Unsupported format: {target_format}")

        return {"filename": output_path.name}

    def _parse_table(self, text, filename):
        rows = []
        for line in (text or "").splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if "|" in stripped:
                cells = [cell.strip() for cell in stripped.split("|") if cell.strip()]
            elif "," in stripped:
                cells = [cell.strip() for cell in stripped.split(",")]
            else:
                continue
            if len(cells) >= 2:
                rows.append(cells)
        return {
            "parser_type": "table",
            "filename": filename,
            "rows": rows[:100],
            "row_count": len(rows),
        }

    def _parse_resume(self, text, filename):
        prompt = (
            f"Extract resume details from the following text.\n"
            f"Return JSON with keys: name, headline, skills, experience, education, contact.\n\n{text[:12000]}"
        )
        result = get_client().chat_once(
            messages=[{"role": "user", "content": prompt}],
            options={"max_new_tokens": 500, "temperature": 0.1, "do_sample": False},
        )
        return self._json_or_text("resume", filename, result.get("message", {}).get("content", ""))

    def _parse_invoice(self, text, filename):
        prompt = (
            f"Extract invoice or receipt data from the following text.\n"
            f"Return JSON with keys: vendor, invoice_number, invoice_date, due_date, currency, totals, line_items.\n\n{text[:12000]}"
        )
        result = get_client().chat_once(
            messages=[{"role": "user", "content": prompt}],
            options={"max_new_tokens": 500, "temperature": 0.1, "do_sample": False},
        )
        return self._json_or_text("invoice", filename, result.get("message", {}).get("content", ""))

    def _parse_contract(self, text, filename):
        prompt = (
            f"Review the following contract text.\n"
            "Return JSON with keys: parties, effective_date, term, payment_terms, termination, risks, missing_items.\n\n"
            f"{text[:14000]}"
        )
        result = get_client().chat_once(
            messages=[{"role": "user", "content": prompt}],
            options={"max_new_tokens": 700, "temperature": 0.15, "do_sample": False},
        )
        return self._json_or_text("contract", filename, result.get("message", {}).get("content", ""))

    def _parse_sections(self, text, filename):
        chunks = self._chunk_text(text, chunk_size=3000, overlap=150)
        sections = []
        for idx, chunk in enumerate(chunks[:8], start=1):
            prompt = (
                f"Summarize section {idx} of file {filename}.\n"
                "Return a short heading and 2-4 bullet summary.\n\n"
                f"{chunk}"
            )
            result = get_client().chat_once(
                messages=[{"role": "user", "content": prompt}],
                options={"max_new_tokens": 180, "temperature": 0.15, "do_sample": False},
            )
            sections.append({"section": idx, "summary": (result.get("message", {}).get("content", "") or "").strip()})
        return {"parser_type": "section-summary", "filename": filename, "sections": sections}

    @staticmethod
    def _json_or_text(parser_type, filename, content):
        text = (content or "").strip()
        try:
            return {
                "parser_type": parser_type,
                "filename": filename,
                "data": json.loads(text),
            }
        except json.JSONDecodeError:
            return {
                "parser_type": parser_type,
                "filename": filename,
                "raw": text,
            }
